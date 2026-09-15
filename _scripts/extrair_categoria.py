#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extrai dados dos comprovantes de uma categoria e gera, DENTRO da pasta:
  - dados_extraidos.json   (dados organizados por comprovante)
  - resumo_extraidos.xlsx  (dados tabulados)

Uso: python extrair_categoria.py "Comunicação e tecnologia"

Campos extraídos por comprovante: data, valor, pagador (quem pagou),
recebedor (quem recebeu), categoria e tipo de comprovante.
"""
import os
import re
import sys
import json
import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import BASE, COMPROVANTES_DIR, CONFIG, DADOS  # noqa: E402
from indices_categoria import listar_comprovantes, pasta_categoria, rel_base, salvar_indice_categoria  # noqa: E402

MESES = {
    "JAN": 1, "FEV": 2, "MAR": 3, "ABR": 4, "MAI": 5, "JUN": 6,
    "JUL": 7, "AGO": 8, "SET": 9, "OUT": 10, "NOV": 11, "DEZ": 12,
}


def parse_valor(texto):
    """Valor monetário (float) com prioridade para padrões de comprovante."""
    def to_float(s):
        s = s.replace("R$", "").replace(" ", "").strip()
        s = re.sub(r"[^\d,\.]", "", s)
        if not s:
            return None
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        try:
            return round(float(s), 2)
        except ValueError:
            return None

    for pat in [
        r"Valor original\s*R\$\s*([\d.,]+)",
        r"Valor final\s*R\$\s*([\d.,]+)",
        r"Valor\s*R\$\s*([\d.,]+)",
        r"Total a pagar\s*R\$\s*([\d.,]+)",
        r"TOTAL A PAGAR\s*R\$\s*([\d.,]+)",
        r"VALOR TOTAL[:]?\s*([\d.,]+)",
        r"Valor Total[:]?\s*([\d.,]+)",
        r"R\$\s*([\d.,]+)",
    ]:
        m = re.search(pat, texto, re.IGNORECASE)
        if m:
            v = to_float(m.group(1))
            if v is not None and v > 0:
                return v
    return None


def parse_data(texto):
    """Data do comprovante -> ISO (YYYY-MM-DD)."""
    m = re.search(
        r"(?<![A-Za-z0-9])([0-9Oo]{1,2})\s*(JAN|FEV|MAR|ABR|MAI|JUN|JUL|AGO|SET|OUT|NOV|DEZ)[.]?\s+(\d{4})(?!\d)",
        texto, re.IGNORECASE,
    )
    if m:
        d, mes, y = int(m.group(1).upper().replace("O", "0")), MESES[m.group(2).upper()], int(m.group(3))
        try:
            return datetime.date(y, mes, d).isoformat()
        except ValueError:
            return None
    m = re.search(r"(?:EMISS[AÃ]O|emissão)[:\s]*(\d{1,2})/(\d{1,2})/(\d{4})", texto, re.IGNORECASE)
    if m:
        return f"{int(m.group(3)):04d}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", texto)
    if m:
        return f"{int(m.group(3)):04d}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    m = re.search(
        r"(\d{1,2})\s+de\s+(janeiro|fevereiro|mar[çc]o|abril|maio|junho|julho|agosto|setembro|outubro|novembro|dezembro)\s+de\s+(\d{4})",
        texto, re.IGNORECASE,
    )
    if m:
        mes_n = {"janeiro": 1, "fevereiro": 2, "março": 3, "marco": 3, "abril": 4,
                 "maio": 5, "junho": 6, "julho": 7, "agosto": 8, "setembro": 9,
                 "outubro": 10, "novembro": 11, "dezembro": 12}
        mes = mes_n.get(m.group(2).lower())
        if mes:
            return f"{int(m.group(3)):04d}-{mes:02d}-{int(m.group(1)):02d}"
    return None


def limpar(s):
    s = s.replace("|", " ").replace("  ", " ")
    s = re.sub(r"\s+", " ", s).strip(" -–.")
    return s.strip()


def limpar_estabelecimento(s):
    """Limpa ruído de OCR comum em 'Estabelecimento' (débito cartão)."""
    s = limpar(s)
    s = re.sub(r"^Estabele\w*\s*", "", s, flags=re.I)
    s = s.replace("cs jpeRMERCADOS", "SUPERMERCADOS")
    s = s.replace("cs jpeRMERCADO", "SUPERMERCADO")
    s = s.replace("jpeRMERCADOS", "SUPERMERCADOS")
    s = s.replace("jpeRMERCADO", "SUPERMERCADO")
    s = s.replace("cs jpe", "SUPER")
    s = s.replace("cs SUPER", "SUPER")
    s = s.replace("«N", "S")
    s = s.replace("MP *", "")
    s = s.replace("MP ", "")
    # correções específicas validadas
    s = s.replace("SpTIPLUS", "SUPTPLUS")
    s = s.replace("sAcOLAO", "SACOLÃO")
    s = s.replace("pAnARIA", "PADARIA")
    s = s.replace("PAO DE MEL", "PÃO DE MEL")
    s = s.replace("AcougueCostello", "Açougue Costello")
    return limpar(s)


RESIDUOS_ESTAB = {"nto", "mento", "ento", "to", "imento", "cimen", "ecimen", "cimento", "ecimento", "uae"}



def eh_comprovante_debito(texto):
    """Detecta comprovantes Nubank de debito mesmo com acento/OCR corrompido."""
    t = texto.lower()
    if re.search(r"compra\s+no\s+d\S{0,4}bito|compra\s+no\s+debito", t, re.IGNORECASE):
        return True
    return "dados da transa" in t and "data e hora" in t and "nsu" in t

def recebedor_debito(texto):
    """Estabelecimento de um comprovante de débito (bloco 'Data e hora'..'NSU')."""
    m = re.search(r"Data e hora[^\n]*\n(.*?)NSU", texto, re.DOTALL)
    if m:
        linhas = [ln.strip() for ln in m.group(1).splitlines() if ln.strip()]
        frags = []
        for ln in linhas:
            ln2 = re.sub(r"^(estabel|ecimen|ecim|to|uae|estabele\w*|cimen\w*)\b", "", ln, flags=re.I).strip()
            if not ln2:
                continue
            if ln2.lower() in RESIDUOS_ESTAB:
                continue
            frags.append(ln2)
        if frags:
            return limpar_estabelecimento(" ".join(frags))
    m = re.search(r"Estabele\w*\s*[:\s]*(.+)", texto, re.IGNORECASE)
    if m:
        return limpar_estabelecimento(m.group(1))
    return None


def recebedor_credito(texto):
    """Estabelecimento de um pagamento no cartão de crédito (bloco 'Destino'..'Origem')."""
    m = re.search(r"Destino\s*\n([\s\S]*?)Origem", texto, re.DOTALL)
    if not m:
        m = re.search(r"Destino\s*\n([\s\S]*?)$", texto, re.DOTALL)
    if not m:
        return None
    frags = []
    for ln in m.group(1).splitlines():
        s = ln.strip()
        if not s:
            continue
        s = re.sub(r"^(estabele\w*|estabel\w*|ecimen\w*|ecim\w*)\b", "", s, flags=re.I).strip()
        # remove prefixo não-alfabético (números, símbolos, espaços)
        s = re.sub(r"^[^A-Za-zÀ-ú]+", "", s).strip()
        if s and s.lower() not in RESIDUOS_ESTAB:
            frags.append(s)
    nome = " ".join(frags)
    nome = re.sub(r"[«»“”|°º]", " ", nome).replace('"', ' ')
    nome = nome.replace("DI*ÍGOOGLE", "GOOGLE").replace("ÍGOOGLE", "GOOGLE")
    nome = nome.replace("o0GLE", "GOOGLE").replace("0GLE", "GOOGLE")
    nome = nome.replace("Spotify Mus", "Spotify Music")
    nome = nome.replace("IFD*IFOOD", "IFOOD")
    return limpar(nome)


# =================== EXTRATORES POR CATEGORIA ===================

def extrair_alimentacao(it):
    texto = it["texto_extraido"]
    t = texto.lower()
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Alimentação",
        "tipo": None,
        "data": parse_data(texto),
        "valor": parse_valor(texto),
        "pagador": None,
        "recebedor": None,
    }

    if "danfe" in t or "nf-e" in t:
        r["tipo"] = "DANFE/NF-e"
        m = re.search(r"RECEBEMOS\s+(.+)", texto, re.IGNORECASE)
        if m:
            r["recebedor"] = limpar(m.group(1).split("OS PRODUTOS")[0].split("-")[0])
        m = re.search(r"DEST\.?:?\s*(.+)", texto, re.IGNORECASE)
        if m:
            r["pagador"] = limpar(m.group(1).split("-")[0])
    elif eh_comprovante_debito(texto):
        r["tipo"] = "Débito (cartão)"
        r["recebedor"] = recebedor_debito(texto)
        m = re.search(
            r"titular do cart[ãa]o\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CPF|N[úu]mero))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    else:
        if "transferência" in t or "transferencia" in t:
            r["tipo"] = "Pix (transferência)"
        else:
            r["tipo"] = "Pix (pagamento)"
        m = re.search(
            r"Origem\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:Institui|CPF|Banco|Conta|Ag|CNPJ))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
        m = re.search(
            r"Destino\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CNPJ|CPF|Institui|Chave|Banco|Tipo))",
            texto, re.IGNORECASE,
        )
        if m:
            r["recebedor"] = limpar(" ".join(m.group(1).split()))
        if not r["pagador"]:
            m = re.search(r"Pagad\s*or\s*\n?\s*(.+)\nBanco", texto, re.IGNORECASE)
            if m:
                r["pagador"] = limpar(m.group(1))
        if not r["recebedor"]:
            m = re.search(r"Favorecido\s*(.+)", texto, re.IGNORECASE)
            if m:
                r["recebedor"] = limpar(m.group(1))
    return r


def extrair_comunicacao(it):
    texto = it["texto_extraido"]
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Comunicação e tecnologia",
        "tipo": "Fatura",
        "data": None,
        "valor": parse_valor(texto),
        "pagador": None,
        "recebedor": None,
    }
    # vencimento: 3ª data do padrão "de X a Y Z"
    m = re.search(
        r"de\s*(\d{1,2}/\d{1,2}/\d{4})\s*a\s*(\d{1,2}/\d{1,2}/\d{4})\s*(\d{1,2}/\d{1,2}/\d{4})",
        texto,
    )
    if m:
        d, mo, y = m.group(3).split("/")
        r["data"] = f"{y}-{mo}-{d}"
    else:
        r["data"] = parse_data(texto)
    # recebedor = operadora (ex.: "Seu número Claro")
    m = re.search(r"Seu n[úu]mero\s+(\w+)", texto, re.IGNORECASE)
    r["recebedor"] = m.group(1).capitalize() if m else "Claro"
    # pagador = titular da linha (maiúsculas antes de "Período de uso")
    m = re.search(r"Vencimento\s*\n\s*([A-ZÀ-Ú][A-ZÀ-Ú\s]{4,}?)\s*Per[íi]odo", texto)
    if m:
        r["pagador"] = limpar(" ".join(m.group(1).split()))
    if not r["pagador"]:
        m = re.search(r"([A-ZÀ-Ú][A-ZÀ-Ú\s]{6,}?)\s*Per[íi]odo de uso", texto)
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    return r


def extrair_lazer(it):
    texto = it["texto_extraido"]
    t = texto.lower()
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Lazer e convivência",
        "tipo": None,
        "data": parse_data(texto),
        "valor": parse_valor(texto),
        "pagador": None,
        "recebedor": None,
    }
    if "compra no débito" in t or "compra no debito" in t:
        r["tipo"] = "Débito (cartão)"
        r["recebedor"] = recebedor_debito(texto)
        m = re.search(
            r"titular do cart[ãa]o\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CPF|N[úu]mero))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    elif "transferência crédito" in t or "transferencia credito" in t:
        r["tipo"] = "Cartão de crédito"
        r["recebedor"] = recebedor_credito(texto)
        m = re.search(
            r"Origem\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:Institui|CPF|Banco|Conta|Ag|CNPJ))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    else:
        r["tipo"] = "Pix (transferência)"
        m = re.search(
            r"Origem\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:Institui|CPF|Banco|Conta|Ag|CNPJ))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
        m = re.search(
            r"Destino\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CNPJ|CPF|Institui|Chave|Banco|Tipo))",
            texto, re.IGNORECASE,
        )
        if m:
            r["recebedor"] = limpar(" ".join(m.group(1).split()))
    return r


def moeda_para_float(s):
    s = s.replace("R$", "").replace(" ", "").strip()
    s = re.sub(r"[^\d,\.]", "", s)
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def iso_ddmm(s):
    """'dd/mm/aaaa' -> 'aaaa-mm-dd'."""
    try:
        d, mo, y = s.strip().split("/")
        return f"{y}-{mo}-{d}"
    except Exception:
        return None


MES_PT = {
    "janeiro": "01", "fevereiro": "02", "março": "03", "marco": "03", "abril": "04",
    "maio": "05", "junho": "06", "julho": "07", "agosto": "08", "setembro": "09",
    "outubro": "10", "novembro": "11", "dezembro": "12",
}

MES_ARQUIVO = {
    "jan": "01", "janeiro": "01",
    "fev": "02", "fevereiro": "02",
    "mar": "03", "marco": "03", "março": "03",
    "abr": "04", "abril": "04",
    "mai": "05", "maio": "05",
    "jun": "06", "junho": "06",
    "jul": "07", "julho": "07",
    "ago": "08", "agosto": "08",
    "set": "09", "setembro": "09",
    "out": "10", "outubro": "10",
    "nov": "11", "novembro": "11",
    "dez": "12", "dezembro": "12",
}


def data_por_nome_arquivo(nome, ano="2026"):
    base = os.path.splitext(os.path.basename(nome or ""))[0].lower()
    base = re.sub(r"[^a-záéíóúâêôãõç0-9]+", " ", base)
    for token in base.split():
        mes = MES_ARQUIVO.get(token)
        if mes:
            return f"{ano}-{mes}-01"
    return None


def extrair_moradia(it):
    texto = it["texto_extraido"]
    t = texto.lower()
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Moradia",
        "tipo": None,
        "data": parse_data(texto),
        "valor": parse_valor(texto),
        "pagador": "Daisy Ferreira Braga de Souza",
        "recebedor": None,
    }

    if "pacto administradora" in t:
        # Boleto de condomínio
        r["tipo"] = "Boleto (condomínio)"
        r["recebedor"] = "Pacto Administradora"
        m = re.search(r"RATEIO MENSAL\s*R\$\s*([\d.,]+)", texto, re.IGNORECASE)
        if m:
            r["valor"] = moeda_para_float(m.group(1))
        m = re.search(r"VENCIMENTO ORIGINAL\s*(\d{1,2}/\d{1,2}/\d{4})", texto)
        if m:
            r["data"] = iso_ddmm(m.group(1))
    elif "recibo do pagador" in t:
        # Boleto (internet)
        r["tipo"] = "Boleto (internet)"
        m = re.search(r"Benefici[áa]rio\s+([^/]+)", texto, re.IGNORECASE)
        if m:
            r["recebedor"] = limpar(m.group(1))
        m = re.search(r"Pagador\s+([^/]+)", texto)
        if m:
            r["pagador"] = limpar(m.group(1))
        m = re.search(r"Vencimento[:\s]*(\d{1,2}/\d{1,2}/\d{4})", texto)
        if m:
            r["data"] = iso_ddmm(m.group(1))
        m = re.search(r"Valor do Documento\s*R\$\s*([\d.,]+)", texto)
        if m:
            r["valor"] = moeda_para_float(m.group(1))
    elif "demonstrativo individual" in t or "consumo de água" in t:
        # Conta de água do condomínio
        r["tipo"] = "Conta de água"
        r["recebedor"] = "Giardino di Napoli (condomínio)"
        m = re.search(r"Valor a pagar\s*R\$\s*([\d.,]+)", texto, re.IGNORECASE)
        if m:
            r["valor"] = moeda_para_float(m.group(1))
        m = re.search(r"Vencimento em\s+([A-Za-zçÇ]+)/(\d{4})", texto, re.IGNORECASE)
        if m:
            mes = MES_PT.get(m.group(1).lower())
            r["data"] = f"{m.group(2)}-{mes}-01" if mes else r["data"]
    elif "unid. consumidora" in t or "metha" in t:
        # Fatura de energia
        r["tipo"] = "Fatura de energia"
        r["recebedor"] = "Metha Energia"
        m = re.search(r"Vencimento:\s*(\d{1,2}/\d{1,2}/\d{4})", texto)
        if m:
            r["data"] = iso_ddmm(m.group(1))
        m = re.search(r"Total a Pagar:?\s*R\$\s*([\d.,]+)", texto, re.IGNORECASE)
        if m:
            r["valor"] = moeda_para_float(m.group(1))
    elif "favorecido" in t:
        # Pix pagamento (IPTU)
        r["tipo"] = "Pix (pagamento)"
        m = re.search(r"Favorecido\s+([^\n]+)", texto, re.IGNORECASE)
        if m:
            r["recebedor"] = limpar(m.group(1))
    elif "preço do gás" in t or ("pedido" in t and "entregue" in t):
        # Gás via app
        r["tipo"] = "Gás (app)"
        m = re.search(r"revenda\s*\n\s*([^\n]+)", texto, re.IGNORECASE)
        if m:
            r["recebedor"] = limpar(m.group(1).split("-")[0])
    elif "transferência" in t or "transferencia" in t:
        r["tipo"] = "Pix (transferência)"
        m = re.search(
            r"Destino\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CNPJ|CPF|Institui|Chave|Banco))",
            texto, re.IGNORECASE,
        )
        if m:
            r["recebedor"] = limpar(" ".join(m.group(1).split()))
        m = re.search(
            r"Origem\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:Institui|CPF|Banco|Conta|Ag|CNPJ))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    else:
        r["tipo"] = "Pix"
        m = re.search(r"Favorecido\s+([^\n]+)", texto, re.IGNORECASE)
        if m:
            r["recebedor"] = limpar(m.group(1))
    if not r["data"]:
        r["data"] = data_por_nome_arquivo(r["arquivo"])
    return r


def extrair_saude(it):
    texto = it["texto_extraido"]
    t = texto.lower()
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Saúde",
        "tipo": None,
        "data": parse_data(texto),
        "valor": parse_valor(texto),
        "pagador": "Daisy Ferreira Braga de Souza",
        "recebedor": None,
    }
    if "nfs-e" in t or "nfse" in t or "nota fiscal de servi" in t:
        r["tipo"] = "NFS-e (odontológico)"
        r["recebedor"] = "Odontofati LTDA"
        m = re.search(r"Gabriel\s+Arthur\s+Braga\s+Fernandes", texto, re.IGNORECASE)
        if m:
            r["pagador"] = "Gabriel Arthur Braga Fernandes"
        m = re.search(r"Valor Total da NFS-e\s*([\d.,]+)", texto, re.IGNORECASE)
        if m:
            r["valor"] = moeda_para_float(m.group(1))
    elif re.search(r"Dr\(a?\)[.:]?", texto, re.IGNORECASE) and "data:" in t:
        r["tipo"] = "Receita médica"
        r["valor"] = None
        m = re.search(r"Dr\(a?\)[.:]?\s*([^\n]+)", texto, re.IGNORECASE)
        if m:
            r["recebedor"] = limpar(re.sub(r"[^\w\s]", " ", m.group(1)))
    elif eh_comprovante_debito(texto):
        r["tipo"] = "Débito (cartão)"
        r["recebedor"] = recebedor_debito(texto)
        m = re.search(
            r"titular do cart[ãa]o\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CPF|N[úu]mero))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    else:
        r["tipo"] = "Pix"
        m = re.search(
            r"Destino\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CNPJ|CPF|Institui|Chave|Banco|Valor|Tipo))",
            texto, re.IGNORECASE,
        )
        if m:
            r["recebedor"] = limpar(" ".join(m.group(1).split()))
        m = re.search(
            r"Origem\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:Institui|CPF|Banco|Conta|Ag|CNPJ))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    return r


def extrair_transporte(it):
    texto = it["texto_extraido"]
    t = texto.lower()
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Transporte",
        "tipo": "Uber (corrida)",
        "data": parse_data(texto),
        "valor": parse_valor(texto),
        "pagador": "Daisy Ferreira Braga de Souza",
        "recebedor": "Uber",
    }
    m = re.search(
        r"Origem\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:Institui|CPF|Banco|Conta|Ag|CNPJ))",
        texto, re.IGNORECASE,
    )
    if m:
        r["pagador"] = limpar(" ".join(m.group(1).split()))
    if "uber" not in t:
        m = re.search(r"Destino\s*\n\s*(?:Estabele\w*\s*[:\s]*)?(.+)", texto, re.IGNORECASE)
        if m:
            r["recebedor"] = limpar(m.group(1).split("\n")[0])
    return r


def nome_pix(texto, secao):
    """Captura e limpa o nome no bloco 'secao' (Destino/Origem) de um Pix."""
    m = re.search(
        secao + r"\s*\n([\s\S]*?)(?=\n\s*(?:Origem|Destino|Institui|CNPJ|CPF|Chave|Banco|Ag|Conta))",
        texto, re.IGNORECASE,
    )
    if not m:
        return None
    palavras = []
    for ln in m.group(1).splitlines():
        s = ln.strip()
        if not s:
            continue
        if re.match(r"^(Institui|Chave|Ag|Conta|CNPJ|CPF|Tipo)\b", s, re.I):
            continue
        s = re.sub(r"\bNome\b", " ", s, flags=re.I)
        s = re.sub(r"\b[\d.]{3,}\b", " ", s)
        s = re.sub(r"\s+", " ", s).strip()
        if s:
            palavras.append(s)
    return limpar(" ".join(palavras)) if palavras else None


def normalizar_ocr_estendido(texto):
    """Reduz letras repetidas por OCR, ex.: RRRReeeecccceeeebbbbeeeeddddoooorrrr."""
    return re.sub(r"(.)\1{2,}", r"\1", texto)


def nome_bloco_pix_estendido(texto, secao, proxima_secao):
    texto2 = normalizar_ocr_estendido(texto)
    m = re.search(secao + r"\s*\n([\s\S]*?)(?=\n\s*" + proxima_secao + r"|\n\s*CNPJ|\n\s*CPF|$)", texto2, re.IGNORECASE)
    if not m:
        return None
    partes = []
    for ln in m.group(1).splitlines():
        s = ln.strip()
        if not s:
            continue
        s = re.sub(r"\bNome\b", " ", s, flags=re.I)
        s = re.sub(r"\b\d{1,3}(?:\.\d{3})+\b", " ", s)
        s = re.sub(r"\b(?:CNPJ|CPF|Institui[cç][aã]o)\b.*", " ", s, flags=re.I)
        s = re.sub(r"\s+", " ", s).strip()
        if s:
            partes.append(s)
    return limpar(" ".join(partes)) if partes else None


def extrair_educacao(it):
    texto = it["texto_extraido"]
    t = texto.lower()
    t2 = normalizar_ocr_estendido(texto).lower()
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Educação",
        "tipo": None,
        "data": parse_data(texto),
        "valor": parse_valor(texto),
        "pagador": "Daisy Ferreira Braga de Souza",
        "recebedor": None,
    }
    if "inep" in t or "instituto nacional" in t:
        r["tipo"] = "Boleto (ENEM)"
        r["recebedor"] = "INEP (ENEM)"
        m = re.search(r"Valor documento[\s\S]{0,50}?(\d{1,3}(?:\.\d{3})*,\d{2})", texto, re.IGNORECASE)
        if m:
            r["valor"] = moeda_para_float(m.group(1))
        m = re.search(r"Vencimento\s*\n?\s*(\d{1,2}/\d{1,2}/\d{4})", texto)
        if m:
            r["data"] = iso_ddmm(m.group(1))
        m = re.search(r"Sacado\s*\n?\s*CPF:?\s*[\d.\-]+\s*[-–]?\s*([A-ZÀ-Ú][A-ZÀ-Ú ]+)", texto)
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    elif "target" in t:
        r["tipo"] = "Boleto (escola)"
        r["recebedor"] = "Target Educacional Eireli EPP"
        m = re.search(r"Pagador/Avalista[:\s]*([^\n]+)", texto, re.IGNORECASE)
        if m:
            nome = re.sub(r"\([^)]*\)", "", m.group(1))
            nome = nome.replace("DAISYFERREIRABRAGADESOUZA", "DAISY FERREIRA BRAGA DE SOUZA")
            r["pagador"] = limpar(nome)
        m = re.search(r"Valor[:\s]*R\$\s*([\d.,]+)", texto, re.IGNORECASE)
        if m:
            r["valor"] = moeda_para_float(m.group(1))
    elif "resumo do pedido" in t or "vendido por" in t:
        r["tipo"] = "Compra online (livros)"
        m = re.search(r"Vendido por:\s*([^\n]+)", texto, re.IGNORECASE)
        if m:
            r["recebedor"] = limpar(m.group(1))
        m = re.search(r"Total geral[:\s]*R\$\s*([\d.,]+)", texto, re.IGNORECASE)
        if m:
            r["valor"] = moeda_para_float(m.group(1))
    elif "transferência" in t or "transferencia" in t or "pix" in t2:
        r["tipo"] = "Pix (transferência)"
        r["recebedor"] = nome_pix(texto, "Destino")
        if not r["recebedor"]:
            r["recebedor"] = nome_bloco_pix_estendido(texto, "Recebedor", "Pagador")
        m = nome_pix(texto, "Origem")
        if m:
            r["pagador"] = m
        else:
            m = nome_bloco_pix_estendido(texto, "Pagador", "Dados")
            if m:
                r["pagador"] = m
    else:
        r["tipo"] = "Outro"
    return r


def extrair_vestuario(it):
    texto = it["texto_extraido"]
    t = texto.lower()
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Vestuário e higiene",
        "tipo": None,
        "data": parse_data(texto),
        "valor": parse_valor(texto),
        "pagador": "Daisy Ferreira Braga de Souza",
        "recebedor": None,
    }
    if "compra no débito" in t or "compra no debito" in t:
        r["tipo"] = "Débito (cartão)"
        r["recebedor"] = recebedor_debito(texto)
        if r["recebedor"]:
            r["recebedor"] = re.sub(r"^\d+\s*[.\-]?\s*", "", r["recebedor"])
            r["recebedor"] = r["recebedor"].replace("GAS RENNER", "LOJAS RENNER")
            r["recebedor"] = limpar(r["recebedor"])
        m = re.search(
            r"titular do cart[ãa]o\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CPF|N[úu]mero))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    else:
        r["tipo"] = "Pix (transferência)"
        r["recebedor"] = nome_pix(texto, "Destino")
        m = nome_pix(texto, "Origem")
        if m:
            r["pagador"] = m
    return r


def extrair_esporte(it):
    texto = it["texto_extraido"]
    t = texto.lower()
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Esporte e desenvolvimento",
        "tipo": None,
        "data": parse_data(texto),
        "valor": parse_valor(texto),
        "pagador": None,
        "recebedor": None,
    }
    if "compra no débito" in t or "compra no debito" in t:
        r["tipo"] = "Débito (cartão)"
        r["recebedor"] = recebedor_debito(texto)
        m = re.search(
            r"titular do cart[ãa]o\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CPF|N[úu]mero))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    else:
        r["tipo"] = "Pix (transferência)"
        m = re.search(
            r"Origem\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:Institui|CPF|Banco|Conta|Ag|CNPJ))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
        m = re.search(
            r"Destino\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CNPJ|CPF|Institui|Chave|Banco|Tipo))",
            texto, re.IGNORECASE,
        )
        if m:
            r["recebedor"] = limpar(" ".join(m.group(1).split()))
    if not r["data"]:
        r["data"] = data_por_nome_arquivo(r["arquivo"])
    return r


def extrair_mesada(it):
    texto = it["texto_extraido"]
    t = texto.lower()
    r = {
        "arquivo": it["arquivo"],
        "categoria": "Mesada",
        "tipo": None,
        "data": parse_data(texto),
        "valor": parse_valor(texto),
        "pagador": None,
        "recebedor": None,
    }
    if "transferência crédito" in t or "transferencia credito" in t:
        r["tipo"] = "Pix (crédito recebido)"
        m = re.search(
            r"Origem\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:Institui|CPF|Banco|Conta|Ag|CNPJ))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
        m = re.search(
            r"Destino\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:CNPJ|CPF|Institui|Chave|Banco|Tipo))",
            texto, re.IGNORECASE,
        )
        if m:
            r["recebedor"] = limpar(" ".join(m.group(1).split()))
    else:
        r["tipo"] = "Pix"
        m = re.search(
            r"Origem\s*\n\s*Nome\s*([\s\S]*?)(?=\n\s*(?:Institui|CPF|Banco|Conta|Ag|CNPJ))",
            texto, re.IGNORECASE,
        )
        if m:
            r["pagador"] = limpar(" ".join(m.group(1).split()))
    if not r["data"]:
        r["data"] = data_por_nome_arquivo(r["arquivo"])
    return r


EXTRACTORS = {
    "Alimentação": extrair_alimentacao,
    "Comunicação e tecnologia": extrair_comunicacao,
    "Lazer e convivência": extrair_lazer,
    "Moradia": extrair_moradia,
    "Saúde": extrair_saude,
    "Transporte": extrair_transporte,
    "Educação": extrair_educacao,
    "Vestuário e higiene": extrair_vestuario,
    "Esporte e desenvolvimento": extrair_esporte,
    "Mesada": extrair_mesada,
}


def estilo_cabecalho(cell):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def item_vazio(caminho, cat, cfg):
    return {
        "categoria": cat,
        "subpasta": cfg["sub"],
        "arquivo": os.path.basename(caminho),
        "caminho_rel": rel_base(caminho),
        "tipo": os.path.splitext(caminho)[1].lower().lstrip("."),
        "valor": None,
        "competencia": None,
        "competencia_origem": None,
        "data_pagamento": None,
        "estabelecimento": None,
        "texto_extraido": "",
    }


def main():
    cat = sys.argv[1] if len(sys.argv) > 1 else "Alimentação"
    cfg = CONFIG[cat]
    sub = pasta_categoria(cat)
    extractor = EXTRACTORS[cat]

    with open(DADOS, encoding="utf-8") as fh:
        dados = json.load(fh)
    itens = dados["categorias"].get(cat, [])
    por_rel = {os.path.normcase(i.get("caminho_rel", "")): i for i in itens if i.get("caminho_rel")}
    por_nome = {i.get("arquivo", ""): i for i in itens if i.get("arquivo")}

    itens_categoria = []
    for caminho in listar_comprovantes(sub):
        rel = rel_base(caminho)
        item = por_rel.get(os.path.normcase(rel)) or por_nome.get(os.path.basename(caminho))
        if item is None:
            item = item_vazio(caminho, cat, cfg)
        else:
            item = dict(item)
            item["arquivo"] = os.path.basename(caminho)
            item["caminho_rel"] = rel
        itens_categoria.append(item)

    resultados = []
    for it in itens_categoria:
        r = extractor(it)
        caminho_abs = os.path.join(BASE, it.get("caminho_rel", "")) if it.get("caminho_rel") else os.path.join(sub, it["arquivo"])
        r["caminho_rel"] = it.get("caminho_rel") or rel_base(caminho_abs)
        r["caminho"] = caminho_abs
        resultados.append(r)
    resultados.sort(key=lambda r: (r["data"] or "", r["arquivo"]))

    meta = {
        "categoria": cat,
        "pasta": rel_base(sub),
        "total_comprovantes": len(resultados),
        "extraido_em": datetime.datetime.now().isoformat(),
    }

    json_path, xlsx_path, resultados = salvar_indice_categoria(sub, cat, resultados)

    def pct(campo):
        n = sum(1 for r in resultados if r[campo])
        return f"{n}/{len(resultados)}"
    print("Categoria:", cat)
    print("JSON salvo:", json_path)
    print("Planilha salva:", xlsx_path)
    print("Total comprovantes:", len(resultados))
    print(f"Cobertura -> data: {pct('data')} | valor: {pct('valor')} | "
          f"pagador: {pct('pagador')} | recebedor: {pct('recebedor')}")


if __name__ == "__main__":
    main()


