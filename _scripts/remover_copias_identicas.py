#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Remove apenas copias fisicas com MD5 exatamente igual.

Conserva o arquivo ja usado em Controle quando houver; caso contrario, mantem
o primeiro caminho em ordem alfabetica. Um manifesto sem conteudo sensivel e
gravado em _extraidos para auditoria. Nenhum lancamento de planilha e alterado.
"""
import datetime as dt
import hashlib
import json
import os
from collections import defaultdict

from openpyxl import load_workbook

from config import BASE, COMPROVANTES_DIR, EXTRAIDOS, PLANILHA_VBA
from indices_categoria import COMPROVANTE_EXTS, rel_base


def md5(caminho):
    digest = hashlib.md5()
    with open(caminho, "rb") as arquivo:
        for bloco in iter(lambda: arquivo.read(1024 * 1024), b""):
            digest.update(bloco)
    return digest.hexdigest()


def arquivos_referenciados_controle():
    referenciados = set()
    livro = load_workbook(PLANILHA_VBA, read_only=True, data_only=True, keep_vba=True)
    try:
        aba = livro["Controle"]
        for linha in range(4, aba.max_row + 1):
            if not isinstance(aba.cell(linha, 1).value, (int, float)):
                continue
            caminho = aba.cell(linha, 7).value
            if caminho:
                referenciados.add(str(caminho).replace("/", "\\").lower())
    finally:
        livro.close()
    return referenciados


def main():
    por_hash = defaultdict(list)
    for raiz, _pastas, arquivos in os.walk(COMPROVANTES_DIR):
        for nome in arquivos:
            caminho = os.path.join(raiz, nome)
            if os.path.splitext(nome)[1].lower() not in COMPROVANTE_EXTS:
                continue
            por_hash[md5(caminho)].append(caminho)

    usados_no_controle = arquivos_referenciados_controle()
    removidos = []
    for digest, caminhos in sorted(por_hash.items()):
        if len(caminhos) < 2:
            continue
        caminhos.sort(key=lambda item: rel_base(item).lower())
        relacionados = [item for item in caminhos if rel_base(item).replace("/", "\\").lower() in usados_no_controle]
        manter = relacionados[0] if relacionados else caminhos[0]
        for caminho in caminhos:
            if caminho == manter:
                continue
            tamanho = os.path.getsize(caminho)
            os.remove(caminho)
            removidos.append({
                "hash": digest,
                "removido": rel_base(caminho),
                "mantido": rel_base(manter),
                "bytes_removidos": tamanho,
                "criterio": "arquivo referenciado em Controle" if relacionados else "primeiro caminho alfabetico",
            })

    os.makedirs(EXTRAIDOS, exist_ok=True)
    manifesto = {
        "gerado_em": dt.datetime.now().isoformat(timespec="seconds"),
        "regra": "somente arquivos com MD5 identico; um arquivo mantido por grupo",
        "total_removido": len(removidos),
        "bytes_removidos": sum(item["bytes_removidos"] for item in removidos),
        "itens": removidos,
    }
    destino = os.path.join(EXTRAIDOS, "copias_identicas_removidas.json")
    with open(destino, "w", encoding="utf-8") as arquivo:
        json.dump(manifesto, arquivo, ensure_ascii=False, indent=2)
    print(f"Copias removidas: {len(removidos)}")
    print(f"Manifesto: {destino}")


if __name__ == "__main__":
    main()
