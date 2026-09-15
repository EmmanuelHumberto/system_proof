"""Read the saved workbook registry without modifying the workbook."""
from functools import lru_cache
from pathlib import Path
import unicodedata

import openpyxl
from config import PLANILHA_VBA


def chave(value):
    return ''.join(c for c in unicodedata.normalize('NFD', str(value or '').strip().casefold())
                   if not unicodedata.combining(c))


@lru_cache(maxsize=2)
def _read(path, modified, size):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        entries, aliases = {}, {}
        for row in book['Cadastros'].iter_rows(min_row=13, max_col=6, values_only=True):
            category, expense, quota, frequency, active, identifier = row
            if category and expense:
                entries[str(identifier or (category, expense))] = {
                    'categoria': category, 'despesa': expense, 'periodicidade': frequency or 'Mensal',
                    'ativo': chave(active) != 'nao',
                }
        if 'Aliases' in book:
            for category, expense, identifier in book['Aliases'].iter_rows(min_row=2, max_col=3, values_only=True):
                if category and expense and identifier:
                    aliases[(chave(category), chave(expense))] = str(identifier)
        return entries, aliases
    finally:
        book.close()


def catalogo():
    path = Path(PLANILHA_VBA)
    stat = path.stat()
    return _read(str(path), stat.st_mtime_ns, stat.st_size)


def resolver_despesa(category, expense):
    entries, aliases = catalogo()
    key = chave(category), chave(expense)
    for entry in entries.values():
        if (chave(entry['categoria']), chave(entry['despesa'])) == key:
            return entry['despesa']
    entry = entries.get(aliases.get(key))
    if entry and chave(entry['categoria']) == chave(category):
        return entry['despesa']
    return expense


def despesas_ativas(category):
    entries, _ = catalogo()
    return [x['despesa'] for x in entries.values() if x['ativo'] and chave(x['categoria']) == chave(category)]


def periodicidade(category, expense, default='Mensal'):
    entries, _ = catalogo()
    name = resolver_despesa(category, expense)
    for entry in entries.values():
        if chave(entry['categoria']) == chave(category) and chave(entry['despesa']) == chave(name):
            return entry['periodicidade']
    return default
