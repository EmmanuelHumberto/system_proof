#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Utilitarios para manter JSON e XLSX por pasta de categoria."""
import datetime
import hashlib
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from classificacao import classificar_despesa
from config import BASE, CONFIG, COMPROVANTES_DIR

COMPROVANTE_EXTS = {".png", ".jpg", ".jpeg", ".pdf"}


def pasta_categoria(categoria):
    cfg = CONFIG[categoria]
    return os.path.join(COMPROVANTES_DIR, cfg["pasta_top"], cfg["sub"])


def rel_base(caminho):
    if not caminho:
        return ""
    try:
        return os.path.relpath(caminho, BASE)
    except ValueError:
        return caminho


def hash_arquivo(caminho):
    try:
        with open(caminho, "rb") as fh:
            return hashlib.md5(fh.read()).hexdigest()
    except OSError:
        return ""


def listar_comprovantes(pasta):
    itens = []
    if not os.path.isdir(pasta):
        return itens
    for nome in sorted(os.listdir(pasta)):
        caminho = os.path.join(pasta, nome)
        if not os.path.isfile(caminho):
            continue
        if nome.startswith(".~lock"):
            continue
        if os.path.splitext(nome)[1].lower() not in COMPROVANTE_EXTS:
            continue
        itens.append(caminho)
    return itens


def normalizar_resultado(resultado, categoria, caminho_abs=None):
    r = dict(resultado)
    caminho = caminho_abs or r.get("caminho") or ""
    caminho_rel = r.get("caminho_rel") or rel_base(caminho)
    arquivo = r.get("arquivo") or os.path.basename(caminho)
    if caminho and not os.path.isabs(caminho):
        caminho = os.path.join(BASE, caminho)

    r["arquivo"] = os.path.basename(arquivo)
    r["categoria"] = r.get("categoria") or categoria
    r["caminho_rel"] = caminho_rel
    r["hash"] = r.get("hash") or hash_arquivo(caminho)
    r["despesa"] = r.get("despesa") or classificar_despesa(categoria, r)
    data = r.get("data") or ""
    r["mes"] = r.get("mes") or data[:7]
    r["periodicidade"] = r.get("periodicidade") or (
        "Anual" if any(k in (r["despesa"] or "").lower() for k in ("iptu", "enem", "matricula", "matrícula")) else "Mensal"
    )
    return r


def salvar_indice_categoria(pasta, categoria, resultados):
    normalizados = []
    for r in resultados:
        caminho = r.get("caminho") or os.path.join(pasta, r.get("arquivo", ""))
        normalizados.append(normalizar_resultado(r, categoria, caminho))

    meta = {
        "categoria": categoria,
        "pasta": rel_base(pasta),
        "total_comprovantes": len(normalizados),
        "atualizado_em": datetime.datetime.now().isoformat(timespec="seconds"),
        "fonte": "indice da pasta de categoria",
    }

    json_path = os.path.join(pasta, "dados_extraidos.json")
    import json
    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump({**meta, "comprovantes": normalizados}, fh, ensure_ascii=False, indent=2)

    xlsx_path = os.path.join(pasta, "resumo_extraidos.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = categoria[:31]
    cab = [
        "Arquivo", "Hash", "Categoria", "Despesa", "Data", "Mes", "Valor (R$)",
        "Pagador", "Recebedor", "Tipo", "Periodicidade", "Caminho relativo",
    ]
    ws.append(cab)
    for col in range(1, len(cab) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in normalizados:
        ws.append([
            r.get("arquivo") or "",
            r.get("hash") or "",
            r.get("categoria") or categoria,
            r.get("despesa") or "",
            r.get("data") or "",
            r.get("mes") or "",
            r.get("valor") if r.get("valor") is not None else None,
            r.get("pagador") or "",
            r.get("recebedor") or "",
            r.get("tipo") or "",
            r.get("periodicidade") or "Mensal",
            r.get("caminho_rel") or "",
        ])

    larguras = [34, 34, 24, 28, 12, 10, 12, 28, 32, 22, 14, 70]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=7).number_format = '#,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    wb.save(xlsx_path)
    return json_path, xlsx_path, normalizados


def mesclar_indice_categoria(pasta, categoria, novos_resultados):
    """Atualiza o indice da categoria sem perder registros ja existentes."""
    existentes = []
    json_path = os.path.join(pasta, "dados_extraidos.json")
    if os.path.exists(json_path):
        import json
        with open(json_path, encoding="utf-8") as fh:
            existentes = json.load(fh).get("comprovantes", [])

    por_chave = {}
    for r in existentes:
        caminho = os.path.join(pasta, r.get("arquivo", ""))
        nr = normalizar_resultado(r, categoria, caminho)
        chave = nr.get("hash") or nr.get("arquivo")
        if chave:
            por_chave[chave] = nr

    for r in novos_resultados:
        caminho = r.get("caminho") or os.path.join(pasta, r.get("arquivo", ""))
        nr = normalizar_resultado(r, categoria, caminho)
        chave = nr.get("hash") or nr.get("arquivo")
        if chave:
            por_chave[chave] = nr

    for caminho in listar_comprovantes(pasta):
        h = hash_arquivo(caminho)
        if h and h not in por_chave:
            por_chave[h] = normalizar_resultado({
                "arquivo": os.path.basename(caminho),
                "categoria": categoria,
                "caminho_rel": rel_base(caminho),
                "hash": h,
                "tipo": os.path.splitext(caminho)[1].lower().lstrip("."),
                "data": "",
                "valor": None,
                "pagador": "",
                "recebedor": "",
                "periodicidade": "Mensal",
            }, categoria, caminho)

    todos = sorted(por_chave.values(), key=lambda r: (r.get("data") or "", r.get("arquivo") or ""))
    return salvar_indice_categoria(pasta, categoria, todos)
